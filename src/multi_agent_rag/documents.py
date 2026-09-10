"""Document ingestion utilities for local files."""

from __future__ import annotations

import csv
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from multi_agent_rag.models import Document

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
JSON_EXTENSIONS = {".json"}
CSV_EXTENSIONS = {".csv"}
PDF_EXTENSIONS = {".pdf"}
WORD_EXTENSIONS = {".docx"}
PRESENTATION_EXTENSIONS = {".pptx"}
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS | JSON_EXTENSIONS | CSV_EXTENSIONS | PDF_EXTENSIONS | WORD_EXTENSIONS | PRESENTATION_EXTENSIONS
)
PDF_PAGE_BREAK_MARKER = "<!-- rag-page-break -->"
PPTX_SLIDE_BREAK_MARKER = "<!-- rag-slide-break -->"


def load_document(path: str | Path) -> Document:
    """Load a local document into normalized text and metadata."""
    document_path = Path(path)
    if not document_path.exists():
        raise FileNotFoundError(f"Document not found: {document_path}")
    if not document_path.is_file():
        raise ValueError(f"Document path is not a file: {document_path}")

    extension = document_path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        text = document_path.read_text(encoding="utf-8")
        document_type = "text"
    elif extension in JSON_EXTENSIONS:
        text = _read_json(document_path)
        document_type = "json"
    elif extension in CSV_EXTENSIONS:
        text = _read_csv(document_path)
        document_type = "csv"
    elif extension in PDF_EXTENSIONS:
        text, page_count = _read_pdf(document_path)
        document_type = "pdf"
    elif extension in WORD_EXTENSIONS:
        text, word_metadata = _read_docx(document_path)
        document_type = "word"
    elif extension in PRESENTATION_EXTENSIONS:
        text, presentation_metadata = _read_pptx(document_path)
        document_type = "presentation"
    else:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported document extension '{extension}'. Supported extensions: {supported}")

    metadata = {
        "source_path": str(document_path),
        "document_type": document_type,
        "extension": extension,
    }
    if extension in PDF_EXTENSIONS:
        metadata["page_count"] = str(page_count)
    elif extension in WORD_EXTENSIONS:
        metadata.update(word_metadata)
    elif extension in PRESENTATION_EXTENSIONS:
        metadata.update(presentation_metadata)

    return Document(
        title=document_path.name,
        text=_normalize_document_text(text),
        metadata=metadata,
    )


def _read_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8"))
    lines = list(_flatten_json(data))
    return "\n".join(lines)


def _flatten_json(value: Any, prefix: str = "") -> list[str]:
    if isinstance(value, dict):
        lines: list[str] = []
        for key, item in value.items():
            next_prefix = f"{prefix}.{key}" if prefix else str(key)
            lines.extend(_flatten_json(item, next_prefix))
        return lines
    if isinstance(value, list):
        lines = []
        for index, item in enumerate(value):
            next_prefix = f"{prefix}[{index}]" if prefix else f"[{index}]"
            lines.extend(_flatten_json(item, next_prefix))
        return lines
    label = prefix or "value"
    return [f"{label}: {value}"]


def _read_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames:
            rows = []
            for index, row in enumerate(reader, start=1):
                cells = [f"{key}: {value}" for key, value in row.items()]
                rows.append(f"Row {index}: " + "; ".join(cells))
            return "\n".join(rows)

    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        return "\n".join(", ".join(cell for cell in row) for row in reader)


def _read_pdf(path: Path) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF ingestion requires pypdf. Install project dependencies before loading PDF files.") from exc

    reader = PdfReader(str(path))
    pages = [_normalize_extracted_text(page.extract_text() or "") for page in reader.pages]
    if not any(pages):
        raise ValueError(f"No extractable text found in PDF: {path}")
    return f"\n\n{PDF_PAGE_BREAK_MARKER}\n\n".join(pages), len(pages)


def _read_docx(path: Path) -> tuple[str, dict[str, str]]:
    try:
        from docx import Document as WordDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError("DOCX ingestion requires python-docx. Install project dependencies before loading Word files.") from exc

    document = WordDocument(str(path))
    blocks: list[str] = []
    paragraph_count = 0
    table_count = 0

    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = _format_docx_paragraph(paragraph)
            if text:
                blocks.append(text)
                paragraph_count += 1
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            text = _format_docx_table(table)
            if text:
                blocks.append(text)
                table_count += 1

    if not blocks:
        raise ValueError(f"No extractable text found in DOCX: {path}")

    properties = document.core_properties
    metadata = {
        "paragraph_count": str(paragraph_count),
        "table_count": str(table_count),
        "extraction_method": "python-docx",
    }
    if properties.author:
        metadata["author"] = properties.author
    if properties.subject:
        metadata["subject"] = properties.subject
    return "\n\n".join(blocks), metadata


def _format_docx_paragraph(paragraph: Any) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""

    style = paragraph.style
    style_name = getattr(style, "name", "") or ""
    style_id = getattr(style, "style_id", "") or ""
    heading = re.match(r"Heading\s*(\d+)$", style_name, re.IGNORECASE) or re.match(
        r"Heading(\d+)$", style_id, re.IGNORECASE
    )
    if heading:
        level = min(max(int(heading.group(1)), 1), 6)
        return f"{'#' * level} {text}"
    if style_name.lower().startswith("list bullet") or style_id.lower().startswith("listbullet"):
        return f"- {text}"
    if style_name.lower().startswith("list number") or style_id.lower().startswith("listnumber"):
        return f"1. {text}"
    return text


def _format_docx_table(table: Any) -> str:
    rows = [[_escape_markdown_cell(cell.text) for cell in row.cells] for row in table.rows]
    return _format_markdown_table(rows)


def _escape_markdown_cell(value: str) -> str:
    return _normalize_document_text(value).replace("|", "\\|").replace("\n", "<br>")


def _read_pptx(path: Path) -> tuple[str, dict[str, str]]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX ingestion requires python-pptx. Install project dependencies before loading presentations.") from exc

    presentation = Presentation(str(path))
    slides: list[str] = []
    text_slide_count = 0
    table_count = 0

    for slide in presentation.slides:
        blocks: list[str] = []
        title_shape = slide.shapes.title
        if title_shape is not None:
            title = _pptx_shape_text(title_shape)
            if title:
                blocks.append(f"# {title}")

        shapes = sorted(slide.shapes, key=lambda shape: (shape.top, shape.left))
        for shape in shapes:
            if title_shape is not None and shape.element is title_shape.element:
                continue
            if getattr(shape, "has_table", False):
                table = _format_pptx_table(shape.table)
                if table:
                    blocks.append(table)
                    table_count += 1
            elif getattr(shape, "has_text_frame", False):
                text = _pptx_shape_text(shape)
                if text:
                    blocks.append(text)

        slide_text = "\n\n".join(blocks)
        slides.append(slide_text)
        if slide_text:
            text_slide_count += 1

    if not any(slides):
        raise ValueError(f"No extractable text found in PPTX: {path}")

    metadata = {
        "slide_count": str(len(presentation.slides)),
        "text_slide_count": str(text_slide_count),
        "table_count": str(table_count),
        "extraction_method": "python-pptx",
    }
    properties = presentation.core_properties
    if properties.author:
        metadata["author"] = properties.author
    if properties.subject:
        metadata["subject"] = properties.subject
    return f"\n\n{PPTX_SLIDE_BREAK_MARKER}\n\n".join(slides), metadata


def _pptx_shape_text(shape: Any) -> str:
    paragraphs = []
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def _format_pptx_table(table: Any) -> str:
    rows = [[_escape_markdown_cell(cell.text) for cell in row.cells] for row in table.rows]
    return _format_markdown_table(rows)


def _format_markdown_table(rows: list[list[str]]) -> str:
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    lines = [f"| {' | '.join(normalized[0])} |", f"| {' | '.join(['---'] * width)} |"]
    lines.extend(f"| {' | '.join(row)} |" for row in normalized[1:])
    return "\n".join(lines)


def _normalize_extracted_text(text: str) -> str:
    text = _normalize_document_text(text)
    text = re.sub(r"([A-Za-z])-\n([A-Za-z])", r"\1\2", text)
    lines = text.splitlines()
    structured: list[str] = []
    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            if structured and structured[-1]:
                structured.append("")
            continue

        next_line = lines[index + 1].strip() if index + 1 < len(lines) else ""
        if _looks_like_pdf_boundary(stripped, next_line):
            if structured and structured[-1]:
                structured.append("")
            structured.append(stripped)
            continue

        if structured and structured[-1] and _is_pdf_continuation(structured[-1], stripped):
            structured[-1] = f"{structured[-1]} {stripped}"
        else:
            structured.append(stripped)
    return _collapse_blank_lines("\n".join(structured))


def _looks_like_pdf_boundary(line: str, next_line: str) -> bool:
    if line.startswith(("- ", "* ", "|", "#")):
        return False
    if line[0].islower():
        return False
    if re.search(r"\b\d{2}/\d{4}\s*-\s*\d{2}/\d{4}\b", line):
        return True
    next_starts_lowercase = bool(next_line) and next_line[0].islower()
    return len(line) <= 40 and len(line.split()) <= 5 and not next_starts_lowercase and not line.endswith((",", ";", ":"))


def _is_pdf_continuation(previous: str, current: str) -> bool:
    if current.startswith(("- ", "* ", "|", "#")):
        return False
    return current[0].islower() or previous.endswith((",", ";", "-"))


def _normalize_document_text(text: str) -> str:
    replacements = {
        "\ufeff": "",
        "\u00a0": " ",
        "\u200b": "",
        "\u200c": "",
        "\u200d": "",
        "\u2060": "",
        "\uf020": " ",
        "\uf06c": "- ",
        "\uf0b7": "- ",
        "\u2022": "- ",
        "\u25cf": "- ",
        "\ufffd": "",
        "\ufb00": "ff",
        "\ufb01": "fi",
        "\ufb02": "fl",
        "\ufb03": "ffi",
        "\ufb04": "ffl",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)

    text = "".join(_clean_character(character) for character in unicodedata.normalize("NFKC", text))
    return _collapse_blank_lines(text)


def _clean_character(character: str) -> str:
    if character in {"\n", "\t"}:
        return character
    category = unicodedata.category(character)
    if category.startswith("C"):
        return ""
    if category == "Co":
        return ""
    return character


def _collapse_blank_lines(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    normalized_lines: list[str] = []
    blank_seen = False
    for line in lines:
        if not line:
            if not blank_seen:
                normalized_lines.append("")
            blank_seen = True
            continue
        normalized_lines.append(line)
        blank_seen = False
    return "\n".join(normalized_lines).strip()
