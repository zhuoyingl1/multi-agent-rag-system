import pytest
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock

pytest.importorskip("fastapi")
pytest.importorskip("anyio")

from fastapi.testclient import TestClient
from docx import Document as WordDocument

from multi_agent_rag.api.main import build_app, run_query
from multi_agent_rag.ingestion import (
    CHUNKING_VERSION,
    INDEX_VERSION,
    DocumentBusyError,
    DocumentCleanupError,
    RegisteredDocument,
)
from multi_agent_rag.persistence import (
    ChunkRecord,
    ConversationMessage,
    ConversationRecord,
    DocumentRecord,
    DocumentStatus,
    KnowledgeSpaceRecord,
)


def fake_document(status: DocumentStatus = DocumentStatus.PROCESSING) -> DocumentRecord:
    now = datetime.now(UTC)
    return DocumentRecord(
        document_id="507f1f77bcf86cd799439011",
        title="uploaded.md",
        file_type="md",
        file_path="output/uploads/uploaded.md",
        file_size=70,
        file_hash="hash",
        status=status,
        progress_percentage=100 if status is DocumentStatus.COMPLETED else 0,
        current_stage=status.value,
        stage_details="",
        created_at=now,
        updated_at=now,
        index_version=INDEX_VERSION if status is DocumentStatus.COMPLETED else None,
        chunking_version=CHUNKING_VERSION if status is DocumentStatus.COMPLETED else None,
        embedding_model="nomic-embed-text" if status is DocumentStatus.COMPLETED else None,
        chunk_count=2 if status is DocumentStatus.COMPLETED else 0,
    )


def fake_knowledge_space() -> KnowledgeSpaceRecord:
    now = datetime.now(UTC)
    return KnowledgeSpaceRecord(
        knowledge_space_id="507f1f77bcf86cd799439012",
        name="Research",
        description="Related research documents",
        created_at=now,
        updated_at=now,
    )


def test_health_endpoint() -> None:
    client = TestClient(build_app())

    response = client.get("/health")

    assert response.status_code == 200
    assert response.json()["status"] == "healthy"


def test_cors_allows_local_frontend() -> None:
    client = TestClient(build_app())

    response = client.options(
        "/query",
        headers={
            "Origin": "http://localhost:3000",
            "Access-Control-Request-Method": "POST",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://localhost:3000"


def test_cors_allows_next_fallback_port() -> None:
    client = TestClient(build_app())

    response = client.options(
        "/query",
        headers={
            "Origin": "http://localhost:3001",
            "Access-Control-Request-Method": "POST",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://localhost:3001"


def test_knowledge_space_endpoints_create_and_list(monkeypatch) -> None:
    spaces = MagicMock()
    spaces.create.return_value = fake_knowledge_space()
    spaces.list.return_value = [fake_knowledge_space()]
    spaces.count.return_value = 1
    documents = MagicMock()
    documents.count.return_value = 2
    monkeypatch.setattr("multi_agent_rag.api.main.KnowledgeSpaceRepository.from_store", lambda _store: spaces)
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: documents)
    client = TestClient(build_app())

    created = client.post(
        "/knowledge-spaces",
        json={"name": "Research", "description": "Related research documents"},
    )
    listed = client.get("/knowledge-spaces?limit=10")

    assert created.status_code == 200
    assert created.json()["name"] == "Research"
    assert listed.status_code == 200
    assert listed.json()["knowledge_spaces"][0]["document_count"] == 2
    documents.count.assert_called_once_with(knowledge_space_id="507f1f77bcf86cd799439012")


def test_query_endpoint_returns_grounded_answer() -> None:
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "How does RAG reduce hallucination?",
            "orchestrator": "local",
            "retrieval_backend": "local",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["answer"].startswith("Based on the retrieved evidence")
    assert data["metrics"]["retrieved_sources"] >= 1
    assert data["metrics"]["mode"] == "deterministic_local"
    assert data["workflow_trace"]["mode"] == "deterministic_local"
    assert data["workflow_trace"]["evidence_status"] == "sufficient"
    assert data["workflow_trace"]["selected_agents"]
    assert data["workflow_trace"]["query_intent"] == "general"
    assert data["workflow_trace"]["query_variants"] == 1
    assert data["workflow_trace"]["selected_k"] == len(data["sources"])
    assert data["workflow_trace"]["context_tokens"] > 0
    assert data["workflow_trace"]["selection_reason"] == "fixed"
    assert data["workflow_trace"]["citation_status"] == "missing"
    assert data["citations"]["evidence_count"] == len(data["sources"])
    assert [source["citation_id"] for source in data["sources"]] == [
        f"S{index}" for index in range(1, len(data["sources"]) + 1)
    ]
    assert all(source["source_locator"]["label"] for source in data["sources"])


def test_query_endpoint_uses_llm_answer_by_default(monkeypatch) -> None:
    def fake_post(_self, _path, _payload):
        return {"message": {"content": "RAG reduces hallucination by grounding answers in retrieved evidence [S1]."}}

    monkeypatch.setattr("multi_agent_rag.agents.summarizer.OllamaAnswerComposer._post_json", fake_post)
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={"query": "How does RAG reduce hallucination?", "orchestrator": "local", "retrieval_backend": "local"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == "RAG reduces hallucination by grounding answers in retrieved evidence [S1]."
    assert data["metrics"]["answer_type"] == "llm"
    assert data["metrics"]["answer_model"] == "qwen2.5:3b"
    assert data["citations"]["valid_citation_ids"] == ["S1"]


def test_query_endpoint_uses_persistent_document_id(monkeypatch) -> None:
    expected = run_query(
        "How does RAG reduce hallucination?",
        Path("examples/sample_docs.md"),
        orchestrator="local",
        retrieval_backend="local",
        require_llm_answer=False,
    )
    captured: dict[str, str] = {}

    def fake_document_query(
        query,
        document_id,
        orchestrator,
        require_llm_answer,
        _on_stage,
        _on_answer_delta,
        _retrieval_query,
        _conversation_history,
    ):
        captured["query"] = query
        captured["document_id"] = document_id
        return expected

    monkeypatch.setattr("multi_agent_rag.api.main.safe_run_document_query", fake_document_query)
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "How does RAG reduce hallucination?",
            "document_id": "507f1f77bcf86cd799439011",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 200
    assert captured == {
        "query": "How does RAG reduce hallucination?",
        "document_id": "507f1f77bcf86cd799439011",
    }
    assert response.json()["metrics"]["retrieved_sources"] >= 1


def test_query_endpoint_uses_knowledge_space_scope(monkeypatch) -> None:
    expected = run_query(
        "Compare retrieval approaches.",
        Path("examples/sample_docs.md"),
        orchestrator="local",
        retrieval_backend="local",
        require_llm_answer=False,
    )
    captured: dict[str, str] = {}

    def fake_space_query(query, knowledge_space_id, *_args):
        captured["query"] = query
        captured["knowledge_space_id"] = knowledge_space_id
        return expected

    monkeypatch.setattr("multi_agent_rag.api.main.safe_run_knowledge_space_query", fake_space_query)
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "Compare retrieval approaches.",
            "knowledge_space_id": "507f1f77bcf86cd799439012",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 200
    assert captured == {
        "query": "Compare retrieval approaches.",
        "knowledge_space_id": "507f1f77bcf86cd799439012",
    }


def test_query_endpoint_rejects_multiple_persistent_scopes() -> None:
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "Compare retrieval approaches.",
            "document_id": "507f1f77bcf86cd799439011",
            "knowledge_space_id": "507f1f77bcf86cd799439012",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Choose either document_id or knowledge_space_id, not both."


def test_query_endpoint_rejects_stale_document_index(monkeypatch) -> None:
    documents = MagicMock()
    documents.get.return_value = replace(fake_document(DocumentStatus.COMPLETED), index_version="outdated")
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: documents)
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "How does RAG reduce hallucination?",
            "document_id": "507f1f77bcf86cd799439011",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Document index is stale. Reindex the document before querying it."


def test_create_conversation_scopes_it_to_document(monkeypatch) -> None:
    now = datetime.now(UTC)
    documents = MagicMock()
    documents.get.return_value = fake_document(DocumentStatus.COMPLETED)
    conversations = MagicMock()
    conversations.create.return_value = ConversationRecord(
        conversation_id="conversation-id",
        title="Document review",
        document_id="507f1f77bcf86cd799439011",
        messages=[],
        created_at=now,
        updated_at=now,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: documents)
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.post(
        "/conversations",
        json={"document_id": "507f1f77bcf86cd799439011", "title": "Document review"},
    )

    assert response.status_code == 200
    assert response.json()["conversation_id"] == "conversation-id"
    assert response.json()["document_id"] == "507f1f77bcf86cd799439011"
    conversations.create.assert_called_once_with(
        title="Document review",
        document_id="507f1f77bcf86cd799439011",
    )


def test_create_conversation_scopes_it_to_knowledge_space(monkeypatch) -> None:
    now = datetime.now(UTC)
    conversations = MagicMock()
    conversations.create.return_value = ConversationRecord(
        conversation_id="conversation-id",
        title="Research",
        knowledge_space_id="507f1f77bcf86cd799439012",
        messages=[],
        created_at=now,
        updated_at=now,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.get_ready_knowledge_space", lambda _space_id: fake_knowledge_space())
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.post(
        "/conversations",
        json={"knowledge_space_id": "507f1f77bcf86cd799439012"},
    )

    assert response.status_code == 200
    assert response.json()["knowledge_space_id"] == "507f1f77bcf86cd799439012"
    conversations.create.assert_called_once_with(
        title="Research",
        knowledge_space_id="507f1f77bcf86cd799439012",
    )


def test_list_conversations_returns_scoped_summaries(monkeypatch) -> None:
    now = datetime.now(UTC)
    conversations = MagicMock()
    conversations.list.return_value = [
        ConversationRecord(
            conversation_id="conversation-id",
            title="Document review",
            document_id="507f1f77bcf86cd799439011",
            messages=[ConversationMessage("message-1", "user", "Summarize this.", now)],
            created_at=now,
            updated_at=now,
        )
    ]
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.get(
        "/conversations?document_id=507f1f77bcf86cd799439011&skip=0&limit=20"
    )

    assert response.status_code == 200
    data = response.json()
    assert data["conversations"][0]["title"] == "Document review"
    assert data["conversations"][0]["message_count"] == 1
    conversations.list.assert_called_once_with(
        skip=0,
        limit=20,
        document_id="507f1f77bcf86cd799439011",
        knowledge_space_id=None,
    )


def test_delete_conversation_reports_missing_record(monkeypatch) -> None:
    conversations = MagicMock()
    conversations.delete.return_value = False
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.delete("/conversations/missing-conversation")

    assert response.status_code == 404
    assert response.json()["detail"] == "Conversation not found: missing-conversation"


def test_update_conversation_renames_existing_record(monkeypatch) -> None:
    now = datetime.now(UTC)
    conversations = MagicMock()
    conversations.update_title.return_value = True
    conversations.get.return_value = ConversationRecord(
        conversation_id="conversation-id",
        title="Payment terms",
        document_id="507f1f77bcf86cd799439011",
        messages=[],
        created_at=now,
        updated_at=now,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.put("/conversations/conversation-id", json={"title": "Payment terms"})

    assert response.status_code == 200
    assert response.json()["title"] == "Payment terms"
    conversations.update_title.assert_called_once_with("conversation-id", "Payment terms")


def test_update_conversation_rejects_blank_title(monkeypatch) -> None:
    conversations = MagicMock()
    conversations.update_title.side_effect = ValueError("Conversation title must not be empty.")
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.put("/conversations/conversation-id", json={"title": "   "})

    assert response.status_code == 400
    assert response.json()["detail"] == "Conversation title must not be empty."


def test_query_uses_and_persists_conversation_context(monkeypatch) -> None:
    now = datetime.now(UTC)
    expected = run_query(
        "How does it help?",
        Path("examples/sample_docs.md"),
        orchestrator="local",
        retrieval_backend="local",
        require_llm_answer=False,
    )
    conversations = MagicMock()
    conversations.get.return_value = ConversationRecord(
        conversation_id="conversation-id",
        title="RAG review",
        document_id="507f1f77bcf86cd799439011",
        messages=[
            ConversationMessage("message-1", "user", "What is retrieval augmented generation?", now),
            ConversationMessage("message-2", "assistant", "It grounds answers in evidence [S1].", now),
        ],
        created_at=now,
        updated_at=now,
    )
    captured = {}

    def fake_document_query(*args):
        captured["retrieval_query"] = args[6]
        captured["history"] = args[7]
        return expected

    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    monkeypatch.setattr("multi_agent_rag.api.main.safe_run_document_query", fake_document_query)
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "How does it help?",
            "document_id": "507f1f77bcf86cd799439011",
            "conversation_id": "conversation-id",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 200
    assert "What is retrieval augmented generation?" in captured["retrieval_query"]
    assert len(captured["history"]) == 2
    conversations.add_turn.assert_called_once()
    assert conversations.add_turn.call_args.kwargs["user_content"] == "How does it help?"
    assert conversations.add_turn.call_args.kwargs["assistant_content"] == expected.answer


def test_query_rejects_conversation_from_another_document(monkeypatch) -> None:
    now = datetime.now(UTC)
    conversations = MagicMock()
    conversations.get.return_value = ConversationRecord(
        conversation_id="conversation-id",
        title="Other document",
        document_id="other-document-id",
        messages=[],
        created_at=now,
        updated_at=now,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.ConversationRepository.from_store", lambda _store: conversations)
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "Tell me more.",
            "document_id": "507f1f77bcf86cd799439011",
            "conversation_id": "conversation-id",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Conversation does not belong to the requested document."


def test_query_endpoint_rejects_invalid_orchestrator() -> None:
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "How does RAG reduce hallucination?",
            "orchestrator": "invalid",
            "retrieval_backend": "local",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 422


def test_query_endpoint_rejects_invalid_retrieval_backend() -> None:
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={"query": "How does RAG reduce hallucination?", "retrieval_backend": "invalid", "require_llm_answer": False},
    )

    assert response.status_code == 422


def test_query_endpoint_returns_fallback_for_insufficient_evidence() -> None:
    client = TestClient(build_app())

    response = client.post(
        "/query",
        json={
            "query": "Who won the 1998 world chess championship?",
            "retrieval_backend": "local",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["sources"] == []
    assert data["metrics"]["evidence_status"] == "insufficient"
    assert "No sufficiently relevant retrieved evidence" in data["answer"]


def test_query_endpoint_returns_bad_request_for_missing_document() -> None:
    client = TestClient(build_app())

    response = client.post("/query", json={"query": "What is this?", "document_path": "missing.md"})

    assert response.status_code == 400
    assert "Document not found" in response.json()["detail"]


def test_query_endpoint_rejects_empty_fields() -> None:
    client = TestClient(build_app())

    response = client.post("/query", json={"query": "", "document_path": ""})

    assert response.status_code == 422
    errors = response.json()["detail"]
    assert len(errors) == 2
    assert {tuple(error["loc"]) for error in errors} == {("body", "query"), ("body", "document_path")}


def test_upload_document_returns_queryable_path(monkeypatch) -> None:
    service = MagicMock()
    service.register.side_effect = lambda **values: RegisteredDocument(
        replace(
            fake_document(),
            title=values["title"],
            file_path=values["file_path"],
            file_size=values["file_size"],
            file_hash=values["file_hash"],
            metadata=values["metadata"],
        ),
        duplicate=False,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    upload = client.post(
        "/documents/upload",
        files={"file": ("uploaded.md", b"# Uploaded Notes\n\nRAG uses retrieved evidence to reduce hallucination.", "text/markdown")},
    )

    assert upload.status_code == 200
    uploaded = upload.json()
    assert uploaded["filename"] == "uploaded.md"
    assert uploaded["document_id"] == "507f1f77bcf86cd799439011"
    assert uploaded["document_path"].endswith(".md")
    assert uploaded["size_bytes"] > 0
    assert uploaded["status"] == "processing"
    assert uploaded["duplicate"] is False
    service.process.assert_called_once()

    response = client.post(
        "/query",
        json={
            "query": "How does RAG reduce hallucination?",
            "document_path": uploaded["document_path"],
            "retrieval_backend": "local",
            "require_llm_answer": False,
        },
    )

    assert response.status_code == 200
    assert response.json()["metrics"]["retrieved_sources"] >= 1


def test_upload_document_accepts_docx(monkeypatch) -> None:
    service = MagicMock()
    service.register.side_effect = lambda **values: RegisteredDocument(
        replace(
            fake_document(),
            title=values["title"],
            file_type=values["file_type"],
            file_path=values["file_path"],
            file_size=values["file_size"],
            file_hash=values["file_hash"],
            metadata=values["metadata"],
        ),
        duplicate=False,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())
    document = WordDocument()
    document.add_heading("DOCX Upload", level=1)
    document.add_paragraph("Word documents use the standard ingestion pipeline.")
    content = BytesIO()
    document.save(content)

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "uploaded.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    uploaded = response.json()
    assert uploaded["filename"] == "uploaded.docx"
    assert uploaded["document_path"].endswith(".docx")
    assert service.register.call_args.kwargs["file_type"] == "docx"
    service.process.assert_called_once()


def test_upload_document_rejects_unsupported_extension(monkeypatch) -> None:
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", MagicMock())
    client = TestClient(build_app())

    response = client.post(
        "/documents/upload",
        files={"file": ("notes.exe", b"not a supported document", "application/octet-stream")},
    )

    assert response.status_code == 400
    assert "Unsupported document extension" in response.json()["detail"]


def test_document_status_endpoint_returns_processing_state(monkeypatch) -> None:
    repository = MagicMock()
    repository.get.return_value = fake_document()
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: repository)
    client = TestClient(build_app())

    response = client.get("/documents/507f1f77bcf86cd799439011")

    assert response.status_code == 200
    assert response.json()["status"] == "processing"
    assert response.json()["progress_percentage"] == 0
    assert response.json()["expected_index_version"] == INDEX_VERSION
    assert response.json()["index_stale"] is False


def test_document_progress_stream_emits_changed_and_terminal_states(monkeypatch) -> None:
    repository = MagicMock()
    repository.get.side_effect = [
        fake_document(),
        replace(fake_document(), progress_percentage=20, current_stage="parsing"),
        fake_document(DocumentStatus.COMPLETED),
    ]
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: repository)
    client = TestClient(build_app())

    response = client.get("/documents/507f1f77bcf86cd799439011/progress/stream")

    assert response.status_code == 200
    assert "event: progress" in response.text
    assert '"progress_percentage": 20' in response.text
    assert "event: complete" in response.text
    assert '"status": "completed"' in response.text


def test_document_list_endpoint_returns_paginated_catalog(monkeypatch) -> None:
    repository = MagicMock()
    repository.list.return_value = [fake_document(DocumentStatus.COMPLETED)]
    repository.count.return_value = 1
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: repository)
    client = TestClient(build_app())

    response = client.get("/documents?skip=0&limit=10&status=completed")

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["documents"][0]["filename"] == "uploaded.md"
    assert data["documents"][0]["chunk_count"] == 2
    repository.list.assert_called_once_with(skip=0, limit=10, status=DocumentStatus.COMPLETED)
    repository.count.assert_called_once_with(DocumentStatus.COMPLETED)


def test_document_can_be_assigned_to_knowledge_space(monkeypatch) -> None:
    repository = MagicMock()
    repository.get.side_effect = [
        fake_document(DocumentStatus.COMPLETED),
        replace(
            fake_document(DocumentStatus.COMPLETED),
            knowledge_space_id="507f1f77bcf86cd799439012",
        ),
    ]
    repository.set_knowledge_space.return_value = True
    monkeypatch.setattr("multi_agent_rag.api.main.get_knowledge_space", lambda _space_id: fake_knowledge_space())
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: repository)
    client = TestClient(build_app())

    response = client.put(
        "/documents/507f1f77bcf86cd799439011/knowledge-space",
        json={"knowledge_space_id": "507f1f77bcf86cd799439012"},
    )

    assert response.status_code == 200
    assert response.json()["knowledge_space_id"] == "507f1f77bcf86cd799439012"
    repository.set_knowledge_space.assert_called_once_with(
        "507f1f77bcf86cd799439011",
        "507f1f77bcf86cd799439012",
    )


def test_update_document_renames_catalog_record(monkeypatch) -> None:
    repository = MagicMock()
    repository.update_title.return_value = True
    repository.get.return_value = replace(
        fake_document(DocumentStatus.COMPLETED),
        title="Research notes",
    )
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: repository)
    client = TestClient(build_app())

    response = client.put(
        "/documents/507f1f77bcf86cd799439011",
        json={"title": "Research notes"},
    )

    assert response.status_code == 200
    assert response.json()["filename"] == "Research notes"
    repository.update_title.assert_called_once_with("507f1f77bcf86cd799439011", "Research notes")


def test_document_chunks_endpoint_returns_filtered_locations(monkeypatch) -> None:
    document_repository = MagicMock()
    document_repository.get.return_value = fake_document(DocumentStatus.COMPLETED)
    chunk_repository = MagicMock()
    chunk_repository.list_page.return_value = (
        [
            ChunkRecord(
                chunk_id="chunk-2",
                document_id="507f1f77bcf86cd799439011",
                text="Qdrant stores document vectors.",
                chunk_type="prose",
                index=2,
                metadata={"line_start": "8", "line_end": "9"},
                created_at=datetime.now(UTC),
            )
        ],
        1,
    )
    monkeypatch.setattr("multi_agent_rag.api.main.DocumentRepository.from_store", lambda _store: document_repository)
    monkeypatch.setattr("multi_agent_rag.api.main.ChunkRepository.from_store", lambda _store: chunk_repository)
    client = TestClient(build_app())

    response = client.get(
        "/documents/507f1f77bcf86cd799439011/chunks?skip=0&limit=5&chunk_type=prose&q=Qdrant"
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["chunks"][0]["source_locator"]["label"] == "lines 8-9"
    chunk_repository.list_page.assert_called_once_with(
        "507f1f77bcf86cd799439011",
        skip=0,
        limit=5,
        chunk_type="prose",
        query="Qdrant",
    )


def test_delete_document_endpoint_cleans_registered_document(monkeypatch) -> None:
    service = MagicMock()
    service.delete.return_value = fake_document(DocumentStatus.COMPLETED)
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    response = client.delete("/documents/507f1f77bcf86cd799439011")

    assert response.status_code == 200
    assert response.json() == {
        "document_id": "507f1f77bcf86cd799439011",
        "filename": "uploaded.md",
        "deleted": True,
    }
    service.delete.assert_called_once_with("507f1f77bcf86cd799439011")


def test_delete_document_endpoint_reports_store_failure(monkeypatch) -> None:
    service = MagicMock()
    service.delete.side_effect = DocumentCleanupError("Qdrant unavailable")
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    response = client.delete("/documents/507f1f77bcf86cd799439011")

    assert response.status_code == 503
    assert response.json()["detail"] == "Document cleanup failed: Qdrant unavailable"


def test_reindex_endpoint_queues_existing_document(monkeypatch) -> None:
    service = MagicMock()
    service.prepare_reindex.return_value = fake_document(DocumentStatus.PROCESSING)
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    response = client.post("/documents/507f1f77bcf86cd799439011/reindex")

    assert response.status_code == 200
    assert response.json()["current_stage"] == "processing"
    service.prepare_reindex.assert_called_once_with("507f1f77bcf86cd799439011")
    service.process.assert_called_once()


def test_reindex_endpoint_rejects_concurrent_request(monkeypatch) -> None:
    service = MagicMock()
    service.prepare_reindex.side_effect = DocumentBusyError("Document indexing is already in progress.")
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    response = client.post("/documents/507f1f77bcf86cd799439011/reindex")

    assert response.status_code == 409
    assert "already in progress" in response.json()["detail"]


def test_retry_endpoint_queues_failed_document(monkeypatch) -> None:
    service = MagicMock()
    service.prepare_retry.return_value = fake_document(DocumentStatus.PROCESSING)
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    response = client.post("/documents/507f1f77bcf86cd799439011/retry")

    assert response.status_code == 200
    assert response.json()["status"] == "processing"
    service.prepare_retry.assert_called_once_with("507f1f77bcf86cd799439011")
    service.process.assert_called_once()


def test_retry_endpoint_rejects_non_failed_document(monkeypatch) -> None:
    service = MagicMock()
    service.prepare_retry.side_effect = ValueError("Only failed document indexing can be retried.")
    monkeypatch.setattr("multi_agent_rag.api.main.create_document_ingestion_service", lambda: service)
    client = TestClient(build_app())

    response = client.post("/documents/507f1f77bcf86cd799439011/retry")

    assert response.status_code == 400
    assert response.json()["detail"] == "Only failed document indexing can be retried."


def test_stream_endpoint_returns_all_events() -> None:
    client = TestClient(build_app())

    response = client.post(
        "/query/stream",
        json={"query": "How does RAG reduce hallucination?", "retrieval_backend": "local", "require_llm_answer": False},
    )

    assert response.status_code == 200
    body = response.text
    assert "event: planning" in body
    assert "event: retrieval" in body
    assert "event: agents" in body
    assert "event: judge" in body
    assert "event: answer_delta" in body
    assert "event: final" in body
    assert body.index("event: planning") < body.index("event: retrieval")
    assert body.index("event: retrieval") < body.index("event: agents")
    assert body.index("event: agents") < body.index("event: judge")
    assert body.index("event: answer_delta") < body.index("event: final")


def test_stream_endpoint_forwards_ollama_tokens(monkeypatch) -> None:
    def fake_stream(_self, _path, payload):
        assert payload["stream"] is True
        yield {"message": {"content": "RAG grounds "}, "done": False}
        yield {"message": {"content": "answers in evidence "}, "done": False}
        yield {"message": {"content": "[S1]."}, "done": True}

    monkeypatch.setattr("multi_agent_rag.agents.summarizer.OllamaAnswerComposer._stream_json", fake_stream)
    client = TestClient(build_app())

    response = client.post(
        "/query/stream",
        json={"query": "How does RAG reduce hallucination?", "orchestrator": "local", "retrieval_backend": "local"},
    )

    assert response.status_code == 200
    body = response.text
    assert body.count("event: answer_delta") == 3
    assert 'data: {"delta": "RAG grounds "}' in body
    assert 'data: {"delta": "answers in evidence "}' in body
    assert 'data: {"delta": "[S1]."}' in body
    assert '"answer": "RAG grounds answers in evidence [S1]."' in body
    assert '"streaming_mode": "token"' in body
    assert '"time_to_first_token_ms":' in body


def test_stream_endpoint_returns_bad_request_for_missing_document() -> None:
    client = TestClient(build_app())

    response = client.post("/query/stream", json={"query": "What is this?", "document_path": "missing.md"})

    assert response.status_code == 400
    assert "Document not found" in response.json()["detail"]


def test_metrics_endpoint_updates_after_query() -> None:
    client = TestClient(build_app())

    client.post(
        "/query",
        json={"query": "How does RAG reduce hallucination?", "retrieval_backend": "local", "require_llm_answer": False},
    )
    response = client.get("/health/metrics")

    assert response.status_code == 200
    assert response.json()["run_count"] >= 1


def test_integrations_endpoint_returns_readiness(monkeypatch) -> None:
    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, _exc_type, _exc, _traceback):
            return False

        def read(self):
            return b'{"models":[{"name":"qwen2.5:3b"}]}'

    monkeypatch.setattr("multi_agent_rag.integrations.urlrequest.urlopen", lambda _url, timeout: FakeResponse())
    client = TestClient(build_app())

    response = client.get("/health/integrations")

    assert response.status_code == 200
    data = response.json()
    integration_names = {item["name"] for item in data["integrations"]}
    assert data["integration_count"] == 6
    assert data["integrations"][0]["name"] == "local_hybrid_store"
    assert data["integrations"][0]["status"] == "ready"
    assert "llm_answer" in integration_names


def test_evaluate_endpoint_returns_report() -> None:
    client = TestClient(build_app())

    response = client.post("/evaluate", json={})

    assert response.status_code == 200
    data = response.json()
    assert data["case_count"] == 3
    assert data["passed_count"] == 3
    assert data["failed_count"] == 0


def test_evaluate_endpoint_returns_bad_request_for_missing_cases() -> None:
    client = TestClient(build_app())

    response = client.post("/evaluate", json={"cases_path": "missing-cases.json"})

    assert response.status_code == 400
    assert response.json()["detail"] == "Evaluation cases not found: missing-cases.json"


def test_evaluate_endpoint_rejects_empty_cases_path() -> None:
    client = TestClient(build_app())

    response = client.post("/evaluate", json={"cases_path": ""})

    assert response.status_code == 422
    assert response.json()["detail"][0]["loc"] == ["body", "cases_path"]
