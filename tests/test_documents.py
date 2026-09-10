import json

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from multi_agent_rag.documents import PPTX_SLIDE_BREAK_MARKER, _normalize_document_text, _normalize_extracted_text, load_document
from multi_agent_rag.models import ChunkType
from multi_agent_rag.retrieval.chunking import chunk_document


def test_load_markdown_document(tmp_path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("# Notes\n\nRAG grounds answers in evidence.", encoding="utf-8")

    document = load_document(path)

    assert document.title == "notes.md"
    assert document.metadata["document_type"] == "text"
    assert "RAG grounds" in document.text


def test_load_json_document_flattens_values(tmp_path) -> None:
    path = tmp_path / "metrics.json"
    path.write_text(json.dumps({"metrics": {"grounding": 0.9}, "tags": ["rag", "judge"]}), encoding="utf-8")

    document = load_document(path)

    assert document.metadata["document_type"] == "json"
    assert "metrics.grounding: 0.9" in document.text
    assert "tags[1]: judge" in document.text


def test_load_csv_document_serializes_rows(tmp_path) -> None:
    path = tmp_path / "metrics.csv"
    path.write_text("metric,value\nsource coverage,0.8\nlatency,120\n", encoding="utf-8")

    document = load_document(path)

    assert document.metadata["document_type"] == "csv"
    assert "Row 1: metric: source coverage; value: 0.8" in document.text
    assert "Row 2: metric: latency; value: 120" in document.text


def test_load_document_rejects_unsupported_extension(tmp_path) -> None:
    path = tmp_path / "notes.rtf"
    path.write_text("unsupported", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported document extension"):
        load_document(path)


def test_load_docx_preserves_headings_lists_and_tables(tmp_path) -> None:
    path = tmp_path / "research.docx"
    source = WordDocument()
    source.core_properties.author = "Research Team"
    source.core_properties.subject = "RAG evaluation"
    source.add_heading("System Overview", level=1)
    source.add_paragraph("The pipeline retrieves and reranks grounded evidence.")
    source.add_paragraph("Inspect source citations", style="List Bullet")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Grounding"
    table.cell(1, 1).text = "0.95"
    source.save(path)

    document = load_document(path)
    chunks = chunk_document(document)

    assert document.metadata == {
        "source_path": str(path),
        "document_type": "word",
        "extension": ".docx",
        "paragraph_count": "3",
        "table_count": "1",
        "extraction_method": "python-docx",
        "author": "Research Team",
        "subject": "RAG evaluation",
    }
    assert "# System Overview" in document.text
    assert "- Inspect source citations" in document.text
    assert "| Metric | Value |" in document.text
    assert "| Grounding | 0.95 |" in document.text
    assert ChunkType.TABLE in {chunk.chunk_type for chunk in chunks}


def test_load_docx_rejects_documents_without_extractable_text(tmp_path) -> None:
    path = tmp_path / "empty.docx"
    WordDocument().save(path)

    with pytest.raises(ValueError, match="No extractable text found in DOCX"):
        load_document(path)


def test_load_pptx_preserves_slides_and_tables(tmp_path) -> None:
    path = tmp_path / "research.pptx"
    source = Presentation()
    source.core_properties.author = "Research Team"
    first = source.slides.add_slide(source.slide_layouts[5])
    first.shapes.title.text = "Retrieval Pipeline"
    text_box = first.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    text_box.text_frame.text = "Hybrid retrieval combines keyword, vector, and graph signals."
    table = first.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "Stage"
    table.cell(0, 1).text = "Purpose"
    table.cell(1, 0).text = "Reranking"
    table.cell(1, 1).text = "Improve relevance"
    second = source.slides.add_slide(source.slide_layouts[5])
    second.shapes.title.text = "Evaluation"
    second.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1)).text_frame.text = "Grounding is measured against sources."
    source.save(path)

    document = load_document(path)
    chunks = chunk_document(document)

    assert document.metadata["document_type"] == "presentation"
    assert document.metadata["slide_count"] == "2"
    assert document.metadata["text_slide_count"] == "2"
    assert document.metadata["table_count"] == "1"
    assert document.metadata["extraction_method"] == "python-pptx"
    assert document.metadata["author"] == "Research Team"
    assert "# Retrieval Pipeline" in document.text
    assert "| Reranking | Improve relevance |" in document.text
    assert PPTX_SLIDE_BREAK_MARKER in document.text
    assert {chunk.metadata["slide_start"] for chunk in chunks} == {"1", "2"}
    assert ChunkType.TABLE in {chunk.chunk_type for chunk in chunks}


def test_load_pptx_rejects_presentations_without_extractable_text(tmp_path) -> None:
    path = tmp_path / "empty.pptx"
    Presentation().save(path)

    with pytest.raises(ValueError, match="No extractable text found in PPTX"):
        load_document(path)


def test_load_xlsx_preserves_worksheets_rows_and_formulas(tmp_path) -> None:
    path = tmp_path / "metrics.xlsx"
    source = Workbook()
    summary = source.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Grounding", 0.95])
    details = source.create_sheet("Details")
    details.append(["Component", "Status"])
    details.append(["Retriever", "Ready"])
    details.append(["Total", "=COUNTA(B2:B2)"])
    source.save(path)

    document = load_document(path)
    chunks = chunk_document(document)

    assert document.metadata["document_type"] == "spreadsheet"
    assert document.metadata["worksheet_count"] == "2"
    assert document.metadata["nonempty_worksheet_count"] == "2"
    assert document.metadata["populated_row_count"] == "5"
    assert document.metadata["extraction_method"] == "openpyxl"
    assert "# Worksheet: Summary" in document.text
    assert "| Grounding | 0.95 |" in document.text
    assert "| Total | =COUNTA(B2:B2) |" in document.text
    assert {chunk.metadata.get("section") for chunk in chunks} == {"Worksheet: Summary", "Worksheet: Details"}
    assert sum(chunk.chunk_type is ChunkType.TABLE for chunk in chunks) == 2


def test_load_xlsx_rejects_workbooks_without_extractable_cells(tmp_path) -> None:
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)

    with pytest.raises(ValueError, match="No extractable cells found in XLSX"):
        load_document(path)


def test_normalize_extracted_text_cleans_pdf_artifacts() -> None:
    text = "Skills \uf06c RAG\u00a0systems\n\n\nTools \ufffd FastAPI"

    normalized = _normalize_extracted_text(text)

    assert "\uf06c" not in normalized
    assert "\ufffd" not in normalized
    assert "Skills - RAG systems" in normalized
    assert "\n\n\n" not in normalized


def test_normalize_document_text_removes_common_encoding_artifacts() -> None:
    text = "\ufeffRAG\u200b uses \ufb01ne-grained retrieval.\nPrivate\ue000 marker.\nBad\ufffd char."

    normalized = _normalize_document_text(text)

    assert "\ufeff" not in normalized
    assert "\u200b" not in normalized
    assert "\ue000" not in normalized
    assert "\ufffd" not in normalized
    assert "fine-grained retrieval" in normalized
    assert "Private marker." in normalized


def test_normalize_extracted_text_repairs_pdf_line_wrapping() -> None:
    text = "Retrieval aug-\nmented generation\nuses source evidence.\n\n- Grounding remains visible."

    normalized = _normalize_extracted_text(text)

    assert "augmented generation uses source evidence." in normalized
    assert "- Grounding remains visible." in normalized
    assert "aug-\nmented" not in normalized


def test_normalize_extracted_text_preserves_layout_boundaries() -> None:
    text = (
        "Projects\n"
        "First Project 01/2025-02/2025\n"
        "Tech Stack: Python\n"
        "- Built a service,\n"
        "then deployed it.\n"
        "Second Project 03/2025-04/2025\n"
        "Tech Stack: TypeScript"
    )

    normalized = _normalize_extracted_text(text)

    assert "Projects\n\nFirst Project" in normalized
    assert "deployed it.\n\nSecond Project" in normalized
    assert "- Built a service, then deployed it." in normalized


def test_load_text_document_applies_shared_normalization(tmp_path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("RAG\u00a0systems use \ufb02exible retrieval.\ufffd", encoding="utf-8")

    document = load_document(path)

    assert "RAG systems use flexible retrieval." in document.text
    assert "\ufffd" not in document.text
